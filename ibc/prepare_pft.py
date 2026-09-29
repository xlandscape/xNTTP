"""Prepare an xNTTP IBCgrass PFT workbook from ecotoxicology studies."""

from __future__ import annotations

import argparse
import csv
import math
import os
import statistics
import sys
from pathlib import Path
from typing import Iterable, Sequence

try:
    import openpyxl
except ImportError as error:
    raise SystemExit(
        "openpyxl is required. Run this script with the Python interpreter bundled with xNTTP."
    ) from error


ENDPOINT_COLUMNS = {
    "biomass": ("EC50_biomass", "slope_biomass"),
    "seedling-biomass": ("EC50_SEbiomass", "slope_SEbiomass"),
    "survival": ("EC50_survival", "slope_survival"),
    "establishment": ("EC50_establishment", "slope_establishment"),
    "sterility": ("EC50_sterility", "slope_sterility"),
    "seed-number": ("EC50_seednumber", "slope_seednumber"),
}
SUMMARY_ROWS = {"mean", "sd", "standard deviation"}


def read_dose_response_file(file_path: Path) -> list[tuple[float, float]]:
    """Read species-level EC50 and slope pairs from a tab-separated study file."""
    values: list[tuple[float, float]] = []
    with file_path.open("r", encoding="utf-8-sig", newline="") as input_file:
        reader = csv.reader(input_file, delimiter="\t")
        try:
            header = [value.strip() for value in next(reader)]
        except StopIteration as error:
            raise ValueError(f"{file_path} is empty") from error

        if [value.lower() for value in header[-3:]] != ["spec", "ec50", "b"]:
            raise ValueError(
                f"{file_path} must end its header with the columns Spec, EC50, and b"
            )

        for line_number, row in enumerate(reader, start=2):
            row = [value.strip() for value in row]
            if not row or not any(row):
                continue
            if len(row) < 3:
                raise ValueError(f"{file_path}:{line_number} has fewer than three columns")

            species, ec50_text, slope_text = row[-3:]
            if species.lower() in SUMMARY_ROWS:
                continue
            try:
                ec50 = float(ec50_text)
                slope = float(slope_text)
            except ValueError as error:
                raise ValueError(
                    f"{file_path}:{line_number} contains a non-numeric EC50 or slope"
                ) from error
            if not math.isfinite(ec50) or ec50 <= 0:
                raise ValueError(f"{file_path}:{line_number} EC50 must be finite and positive")
            if not math.isfinite(slope) or slope <= 0:
                raise ValueError(f"{file_path}:{line_number} slope must be finite and positive")
            values.append((ec50, slope))

    if len(values) < 2:
        raise ValueError(f"{file_path} must contain at least two species-level functions")
    return values


def derive_interval(values: Sequence[float], method: str) -> tuple[float, float]:
    """Derive a non-negative interval using the selected method."""
    if method == "range":
        return min(values), max(values)
    if method == "mean-sd":
        mean = statistics.mean(values)
        standard_deviation = statistics.stdev(values)
        return max(0.0, mean - standard_deviation), mean + standard_deviation
    raise ValueError(f"Unsupported interval method: {method}")


def format_interval(interval: tuple[float, float], precision: int) -> str:
    """Format an interval for sampling by the xNTTP IbcGrass component."""
    def format_value(value: float) -> str:
        rounded = f"{value:.{precision}f}".rstrip("0").rstrip(".")
        return rounded if rounded else "0"

    return f"[{format_value(interval[0])};{format_value(interval[1])}]"


def parse_endpoint_specs(specifications: Iterable[str]) -> dict[str, Path]:
    """Parse repeated ENDPOINT=FILE command-line values."""
    endpoints: dict[str, Path] = {}
    for specification in specifications:
        endpoint, separator, file_name = specification.partition("=")
        endpoint = endpoint.strip().lower()
        if not separator or not file_name.strip():
            raise ValueError(f"Invalid endpoint specification: {specification!r}")
        if endpoint not in ENDPOINT_COLUMNS:
            valid = ", ".join(ENDPOINT_COLUMNS)
            raise ValueError(f"Unknown endpoint {endpoint!r}; choose one of: {valid}")
        if endpoint in endpoints:
            raise ValueError(f"Endpoint {endpoint!r} was specified more than once")
        endpoints[endpoint] = Path(file_name.strip())
    return endpoints


def prepare_workbook(
    template_path: Path,
    output_path: Path,
    endpoint_files: dict[str, Path],
    interval_method: str,
    precision: int,
    force: bool,
) -> tuple[dict[str, tuple[str, str]], int]:
    """Create a PFT workbook and return the applied intervals and row count."""
    if not template_path.is_file():
        raise ValueError(f"Template workbook does not exist: {template_path}")
    if template_path.resolve() == output_path.resolve():
        raise ValueError("Output must differ from the template workbook")
    if output_path.exists() and not force:
        raise ValueError(f"Output already exists: {output_path}; use --force to replace it")

    intervals: dict[str, tuple[str, str]] = {}
    for endpoint, file_path in endpoint_files.items():
        if not file_path.is_file():
            raise ValueError(f"Dose-response file does not exist: {file_path}")
        functions = read_dose_response_file(file_path)
        ec50_interval = derive_interval([value[0] for value in functions], interval_method)
        slope_interval = derive_interval([value[1] for value in functions], interval_method)
        intervals[endpoint] = (
            format_interval(ec50_interval, precision),
            format_interval(slope_interval, precision),
        )

    workbook = openpyxl.load_workbook(template_path)
    required_columns = {"ID", "sens"}
    for ec50_column, slope_column in ENDPOINT_COLUMNS.values():
        required_columns.update((ec50_column, slope_column))

    updated_rows = 0
    for worksheet in workbook.worksheets:
        header = {
            cell.value: column_index
            for column_index, cell in enumerate(worksheet[1], start=1)
            if cell.value is not None
        }
        missing = sorted(required_columns.difference(header))
        if missing:
            raise ValueError(
                f"Worksheet {worksheet.title!r} is missing columns: {', '.join(missing)}"
            )

        for row_index in range(2, worksheet.max_row + 1):
            if worksheet.cell(row_index, header["ID"]).value is None:
                continue
            worksheet.cell(row_index, header["sens"]).value = "random"
            for endpoint, (ec50_column, slope_column) in ENDPOINT_COLUMNS.items():
                ec50_value, slope_value = intervals.get(endpoint, (0, 0))
                worksheet.cell(row_index, header[ec50_column]).value = ec50_value
                worksheet.cell(row_index, header[slope_column]).value = slope_value
            updated_rows += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_name(f".{output_path.name}.tmp.xlsx")
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    return intervals, updated_rows


def build_parser() -> argparse.ArgumentParser:
    """Build the command-line parser."""
    parser = argparse.ArgumentParser(
        description="Create an xNTTP IBCgrass PFT workbook from ecotox dose-response files."
    )
    parser.add_argument(
        "--template",
        type=Path,
        default=Path(__file__).with_name("pft.xlsx"),
        help="PFT workbook used as structural and ecological-trait template",
    )
    parser.add_argument("--output", type=Path, required=True, help="New .xlsx file to create")
    parser.add_argument(
        "--endpoint",
        action="append",
        required=True,
        metavar="ENDPOINT=FILE",
        help="Endpoint study input; repeat for each endpoint",
    )
    parser.add_argument(
        "--interval-method",
        choices=("mean-sd", "range"),
        default="mean-sd",
        help="Derive intervals as mean +/- sample SD (default) or observed range",
    )
    parser.add_argument(
        "--precision",
        type=int,
        default=6,
        help="Maximum number of decimal places in generated intervals (default: 6)",
    )
    parser.add_argument("--force", action="store_true", help="Replace an existing output file")
    return parser


def main(arguments: Sequence[str] | None = None) -> int:
    """Run the command-line interface."""
    parser = build_parser()
    options = parser.parse_args(arguments)
    if options.precision < 0 or options.precision > 15:
        parser.error("--precision must be between 0 and 15")
    try:
        endpoint_files = parse_endpoint_specs(options.endpoint)
        intervals, updated_rows = prepare_workbook(
            options.template,
            options.output,
            endpoint_files,
            options.interval_method,
            options.precision,
            options.force,
        )
    except (OSError, ValueError) as error:
        parser.error(str(error))

    print(f"Created {options.output} with {updated_rows} PFT rows set to sens=random")
    for endpoint, (ec50_interval, slope_interval) in intervals.items():
        print(f"{endpoint}: EC50={ec50_interval}, slope={slope_interval}")
    return 0


if __name__ == "__main__":
    sys.exit(main())